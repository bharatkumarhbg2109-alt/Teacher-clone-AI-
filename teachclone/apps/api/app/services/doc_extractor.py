"""Document ingestion — PDF / DOCX / PPTX / TXT / images, with OCR fallback.

Produces page-anchored chunks (like transcription produces time-anchored chunks)
so citations can point at "p.7". Handles scanned PDFs and images via OCR.
"""
import os
from dataclasses import dataclass


@dataclass
class DocChunkData:
    text: str
    page: int | None
    chunk_index: int
    token_count: int
    content_type: str = "document"


_MAX_WORDS = 300
_OVERLAP_WORDS = 30


def _chunk_text(text: str, page: int | None, start_index: int) -> list[DocChunkData]:
    words = text.split()
    chunks: list[DocChunkData] = []
    i = 0
    idx = start_index
    while i < len(words):
        window = words[i : i + _MAX_WORDS]
        if not window:
            break
        chunks.append(
            DocChunkData(
                text=" ".join(window),
                page=page,
                chunk_index=idx,
                token_count=int(len(window) / 0.75),
            )
        )
        idx += 1
        if i + _MAX_WORDS >= len(words):
            break
        i += _MAX_WORDS - _OVERLAP_WORDS
    return chunks


def _ocr_image(image) -> str:
    try:
        import pytesseract

        return pytesseract.image_to_string(image).strip()
    except Exception:
        return ""


class DocExtractor:
    def extract(self, path: str, source_type: str) -> list[DocChunkData]:
        if source_type == "pdf_upload":
            return self.extract_pdf(path)
        if source_type == "image_upload":
            return self.extract_image(path)
        ext = path.rsplit(".", 1)[-1].lower() if "." in path else ""
        if ext == "pdf":
            return self.extract_pdf(path)
        if ext == "docx":
            return self.extract_docx(path)
        if ext == "pptx":
            return self.extract_pptx(path)
        if ext in ("png", "jpg", "jpeg", "webp", "gif", "bmp", "tiff"):
            return self.extract_image(path)
        return self.extract_txt(path)

    # --- PDF ---------------------------------------------------------------
    def extract_pdf(self, path: str) -> list[DocChunkData]:
        import fitz  # PyMuPDF

        chunks: list[DocChunkData] = []
        doc = fitz.open(path)
        idx = 0
        for page_no in range(len(doc)):
            page = doc[page_no]
            text = page.get_text("text").strip()
            # Scanned / image-only page -> OCR the rendered page.
            if len(text) < 20:
                try:
                    from PIL import Image

                    pix = page.get_pixmap(dpi=200)
                    img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                    text = _ocr_image(img)
                except Exception:
                    text = text
            if text:
                page_chunks = _chunk_text(text, page_no + 1, idx)
                chunks.extend(page_chunks)
                idx += len(page_chunks)
        doc.close()
        return chunks

    # --- DOCX --------------------------------------------------------------
    def extract_docx(self, path: str) -> list[DocChunkData]:
        import docx

        d = docx.Document(path)
        text = "\n".join(p.text for p in d.paragraphs if p.text.strip())
        return _chunk_text(text, None, 0)

    # --- PPTX --------------------------------------------------------------
    def extract_pptx(self, path: str) -> list[DocChunkData]:
        from pptx import Presentation

        prs = Presentation(path)
        chunks: list[DocChunkData] = []
        idx = 0
        for slide_no, slide in enumerate(prs.slides, 1):
            parts = [
                shape.text
                for shape in slide.shapes
                if shape.has_text_frame and shape.text.strip()
            ]
            text = "\n".join(parts).strip()
            if text:
                slide_chunks = _chunk_text(text, slide_no, idx)
                chunks.extend(slide_chunks)
                idx += len(slide_chunks)
        return chunks

    # --- TXT / Markdown ----------------------------------------------------
    def extract_txt(self, path: str) -> list[DocChunkData]:
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()
        return _chunk_text(text, None, 0)

    # --- Image -------------------------------------------------------------
    def extract_image(self, path: str) -> list[DocChunkData]:
        from PIL import Image

        img = Image.open(path).convert("RGB")
        text = _ocr_image(img)
        if not text:
            text = f"[Image: {os.path.basename(path)}]"
        return _chunk_text(text, 1, 0)

    def page_count(self, path: str) -> int:
        try:
            import fitz

            doc = fitz.open(path)
            n = len(doc)
            doc.close()
            return n
        except Exception:
            return 0


doc_extractor = DocExtractor()
